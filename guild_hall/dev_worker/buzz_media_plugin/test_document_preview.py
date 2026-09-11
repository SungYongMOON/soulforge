"""Synthetic local rendering checks; no upload or account access."""
import importlib.util
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch, Mock

spec = importlib.util.spec_from_file_location("preview_under_test", Path(__file__).with_name("document_preview.py"))
preview = importlib.util.module_from_spec(spec)
spec.loader.exec_module(preview)


class PreviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.output = self.root / "output"

    def test_real_pdf_first_page_and_bounds(self):
        import pymupdf as fitz
        source = self.root / "source.pdf"
        with fitz.open() as doc:
            doc.new_page(width=2000, height=3000).insert_text((30, 30), "First page")
            doc.new_page().insert_text((30, 30), "Second page")
            doc.save(source)
        original = source.read_bytes()
        result = preview.render_preview(source, self.output)
        self.assertIsNotNone(result)
        pix = fitz.Pixmap(str(result))
        self.assertLessEqual(pix.width, 1000)
        self.assertLessEqual(pix.height, 1400)
        raw = result.read_bytes()
        self.assertEqual(raw[:8], b"\x89PNG\r\n\x1a\n")
        chunks = []
        offset = 8
        while offset < len(raw):
            size = int.from_bytes(raw[offset:offset + 4], "big")
            chunks.append(raw[offset + 4:offset + 8])
            offset += size + 12
        self.assertEqual(offset, len(raw))
        self.assertEqual(set(chunks), {b"IHDR", b"IDAT", b"IEND"})
        self.assertEqual(source.read_bytes(), original)

    def test_bad_pdf_and_unknown_return_none(self):
        for suffix in (".pdf", ".zip", ".unknown"):
            source = self.root / ("source" + suffix)
            source.write_bytes(b"synthetic invalid")
            self.assertIsNone(preview.render_preview(source, self.output))

    def test_timeout_does_not_escape(self):
        source = self.root / "source.pdf"
        source.write_bytes(b"synthetic")
        with patch.object(preview, "_run", side_effect=subprocess.TimeoutExpired("synthetic", 45)):
            self.assertIsNone(preview.render_preview(source, self.output))

    def test_missing_office_is_optional(self):
        source = self.root / "source.docx"
        source.write_bytes(b"synthetic")
        with patch.object(preview, "_libreoffice", return_value=None):
            self.assertIsNone(preview.render_preview(source, self.output))

    def test_child_environment_excludes_profile_credentials(self):
        child = Mock(returncode=0)
        synthetic = {"BUZZ_PRIVATE_KEY": "synthetic", "HERMES_HOME": "synthetic",
                     "API_TOKEN": "synthetic", "PYTHONPATH": "synthetic",
                     "PATH": "synthetic-bin", "TEMP": "synthetic-temp"}
        with patch.dict(preview.os.environ, synthetic, clear=True), patch.object(preview.subprocess, "Popen", return_value=child) as launch:
            preview._run(["synthetic-converter"])
        environment = launch.call_args.kwargs["env"]
        self.assertEqual(set(environment), {"PATH", "TEMP", "PYTHONDONTWRITEBYTECODE"})

    @unittest.skipUnless(preview.os.name == "nt", "Windows tree cleanup")
    def test_timeout_kills_only_owned_windows_tree(self):
        child = Mock(pid=12345, returncode=None)
        child.communicate.side_effect = subprocess.TimeoutExpired("synthetic", 45)
        child.poll.return_value = None
        with patch.object(preview.subprocess, "Popen", return_value=child), patch.object(preview.subprocess, "run") as cleanup:
            with self.assertRaises(subprocess.TimeoutExpired):
                preview._run(["synthetic-converter"])
        self.assertEqual(cleanup.call_args.args[0][1:], ["/PID", "12345", "/T", "/F"])
        child.wait.assert_called_once_with(timeout=10)

    def test_office_uses_isolated_restrictive_profile(self):
        source = self.root / "source.docx"
        source.write_bytes(b"synthetic")
        def run(args):
            self.assertIn("--headless", args)
            profile_arg = next(a for a in args if a.startswith("-env:UserInstallation="))
            self.assertIn("document-preview-", profile_arg)
            local_source = Path(args[-1])
            profile = local_source.parent / "profile" / "user" / "registrymodifications.xcu"
            config = profile.read_text()
            self.assertIn("DisableMacrosExecution", config)
            self.assertIn("BlockUntrustedRefererLinks", config)
            self.assertNotEqual(local_source, source)
        with patch.object(preview, "_libreoffice", return_value="synthetic-soffice"), patch.object(preview, "_run", side_effect=run):
            self.assertIsNone(preview.render_preview(source, self.output))

    @unittest.skipUnless(preview._libreoffice(), "LibreOffice unavailable")
    def test_real_office_first_page(self):
        from openpyxl import Workbook
        source = self.root / "source.xlsx"
        book = Workbook()
        book.active["A1"] = "Synthetic preview"
        book.save(source)
        original = source.read_bytes()
        self.assertIsNotNone(preview.render_preview(source, self.output))
        self.assertEqual(source.read_bytes(), original)

    @unittest.skipUnless(preview._libreoffice(), "LibreOffice unavailable")
    def test_real_pptx_first_slide(self):
        from pptx import Presentation
        source = self.root / "source.pptx"
        deck = Presentation()
        for title in ("Synthetic first slide", "Synthetic second slide"):
            slide = deck.slides.add_slide(deck.slide_layouts[0])
            slide.shapes.title.text = title
        deck.save(source)
        original = source.read_bytes()
        self.assertIsNotNone(preview.render_preview(source, self.output))
        self.assertEqual(source.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
