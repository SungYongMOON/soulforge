# Attachment text worker: one attachment in, one JSON description out.
#
# The bot's model cannot open a .pptx, a .pdf or a .xlsx, and the parsers that
# can are Python libraries already installed on this host. This worker is the
# whole of that bridge: it reads one request on stdin, opens exactly the one file
# the request names, and writes one ASCII JSON object on stdout. It calls no
# model, opens no network, writes no file, and never puts a path or file content
# into an error.
#
# For a presentation it reports, per slide, each shape's identity and box (EMU
# and percent of the slide) beside its text runs, because which side of a drawing
# a number belongs to is a fact about position, not about wording -- a reader who
# only gets the text cannot tell, and a reader who gets the boxes can at least
# see what is adjacent to what. Table cells are reported as cells.
#
# ASCII on the raw buffer: a Korean Windows pipe re-encodes text as cp949, and a
# re-encoded answer is a different answer. The caller runs this with `-I -X utf8`.
import contextlib
import json
import os
import sys

WORKER_SCHEMA = "soulforge.context_attachment_text_worker.v1"
TEXT_FORMATS = ("txt", "md", "csv")
EMU_PER_INCH = 914400


class WorkerError(Exception):
    def __init__(self, code):
        super().__init__(code)
        self.code = code


@contextlib.contextmanager
def _stdout_reserved():
    """Keep file descriptor 1 for the answer and nothing else.

    A parser is free to print: a C extension writes to the process's stdout
    without asking Python, and anything in front of the JSON makes the whole
    answer unreadable to the caller. Descriptor 1 points at stderr while the file
    is being read, where the caller already discards it; the answer is written
    after it is restored.
    """
    sys.stdout.flush()
    saved = os.dup(1)
    os.dup2(2, 1)
    try:
        yield
    finally:
        sys.stdout.flush()
        os.dup2(saved, 1)
        os.close(saved)


def _packages():
    import importlib.metadata as metadata

    found = {}
    for name in ("python-pptx", "pypdf", "PyMuPDF", "openpyxl"):
        try:
            found[name] = metadata.version(name)
        except Exception:
            found[name] = None
    return found


def _percent(value, whole):
    if value is None or not whole:
        return None
    return round(float(value) / float(whole) * 1000.0) / 10.0


def _box(shape, width, height):
    left, top = getattr(shape, "left", None), getattr(shape, "top", None)
    shape_width, shape_height = getattr(shape, "width", None), getattr(shape, "height", None)
    emu = {
        "left": None if left is None else int(left),
        "top": None if top is None else int(top),
        "width": None if shape_width is None else int(shape_width),
        "height": None if shape_height is None else int(shape_height),
    }
    percent = {
        "left": _percent(left, width),
        "top": _percent(top, height),
        "width": _percent(shape_width, width),
        "height": _percent(shape_height, height),
    }
    return {"emu": emu, "percent_of_slide": percent}


def _shape_type(shape):
    try:
        return str(shape.shape_type)
    except Exception:
        return None


def _pptx(path, limits):
    from pptx import Presentation

    presentation = Presentation(path)
    width = int(presentation.slide_width or 0)
    height = int(presentation.slide_height or 0)
    max_shapes = int(limits.get("max_shapes_per_slide", 400))
    max_runs = int(limits.get("max_runs_per_shape", 200))
    slides = []
    for index, slide in enumerate(presentation.slides, start=1):
        shapes = []
        for shape in list(slide.shapes)[:max_shapes]:
            row = {
                "shape_id": int(getattr(shape, "shape_id", 0) or 0),
                "name": str(getattr(shape, "name", "") or ""),
                "shape_type": _shape_type(shape),
                "box": _box(shape, width, height),
                "runs": [],
                "text": "",
                "table": None,
            }
            if getattr(shape, "has_text_frame", False):
                runs = []
                paragraphs = []
                for paragraph in shape.text_frame.paragraphs:
                    line = []
                    for run in paragraph.runs:
                        text = run.text
                        if text is None:
                            continue
                        if len(runs) < max_runs:
                            runs.append(text)
                        line.append(text)
                    paragraphs.append("".join(line))
                row["runs"] = runs
                row["text"] = "\n".join(paragraphs).strip()
            if getattr(shape, "has_table", False):
                table = []
                for table_row in shape.table.rows:
                    table.append([cell.text for cell in table_row.cells])
                row["table"] = table
            shapes.append(row)
        slides.append({
            "slide": index,
            "shapes": shapes,
            "text": "\n".join(row["text"] for row in shapes if row["text"]),
        })
    return {
        "format": "pptx",
        "slide_size": {"emu": {"width": width, "height": height},
                       "inches": {"width": round(width / EMU_PER_INCH, 3) if width else None,
                                  "height": round(height / EMU_PER_INCH, 3) if height else None}},
        "slides": slides,
        "counts": {"slides": len(slides), "shapes": sum(len(slide["shapes"]) for slide in slides)},
    }


def _pdf(path, limits):
    from pypdf import PdfReader

    reader = PdfReader(path)
    max_pages = int(limits.get("max_pages", 200))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        if index > max_pages:
            break
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append({"page": index, "text": text})
    return {"format": "pdf", "pages": pages,
            "counts": {"pages_in_file": len(reader.pages), "pages_read": len(pages)}}


def _xlsx(path, limits):
    from openpyxl import load_workbook

    max_rows = int(limits.get("max_rows_per_sheet", 200))
    max_columns = int(limits.get("max_columns_per_sheet", 40))
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for worksheet in workbook.worksheets:
            cells = []
            for row_index, row in enumerate(worksheet.iter_rows(max_row=max_rows, max_col=max_columns), start=1):
                for cell in row:
                    if cell.value is None:
                        continue
                    cells.append({"ref": str(cell.coordinate), "row": row_index, "value": str(cell.value)})
            sheets.append({"sheet": str(worksheet.title), "cells": cells,
                           "bounds": {"rows": int(worksheet.max_row or 0), "columns": int(worksheet.max_column or 0)},
                           "truncated": bool((worksheet.max_row or 0) > max_rows or (worksheet.max_column or 0) > max_columns)})
        return {"format": "xlsx", "sheets": sheets, "counts": {"sheets": len(sheets)}}
    finally:
        workbook.close()


def _plain(path, fmt, limits):
    max_characters = int(limits.get("max_characters", 200000))
    with open(path, "rb") as handle:
        raw = handle.read(int(limits.get("max_bytes", 8 * 1024 * 1024)))
    for encoding in ("utf-8", "cp949", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            text = None
    if text is None:
        raise WorkerError("unsupported_format")
    return {"format": fmt, "text": text[:max_characters], "characters": len(text),
            "truncated": len(text) > max_characters, "encoding_tried": encoding}


def extract(request):
    path = request.get("path")
    fmt = request.get("format")
    limits = request.get("limits") or {}
    if not isinstance(path, str) or not path or not isinstance(fmt, str):
        raise WorkerError("request_invalid")
    fmt = fmt.lower()
    try:
        with _stdout_reserved():
            if fmt == "pptx":
                body = _pptx(path, limits)
            elif fmt == "pdf":
                body = _pdf(path, limits)
            elif fmt == "xlsx":
                body = _xlsx(path, limits)
            elif fmt in TEXT_FORMATS:
                body = _plain(path, fmt, limits)
            else:
                raise WorkerError("unsupported_format")
    except WorkerError:
        raise
    except ImportError:
        raise WorkerError("parser_unavailable")
    except Exception as error:  # a parser that cannot read this file is not a crash
        return {"status": "error", "code": "parse_failed", "format": fmt,
                "error_type": type(error).__name__, "packages": _packages()}
    return {"status": "ok", "packages": _packages(), **body}


def main():
    raw = sys.stdin.buffer.read(int(64 * 1024 * 1024))
    try:
        request = json.loads(raw.decode("utf-8"))
    except Exception:
        raise WorkerError("request_unreadable")
    if not isinstance(request, dict) or request.get("schema_version") != WORKER_SCHEMA:
        raise WorkerError("request_schema_unknown")
    if request.get("operation") != "extract":
        raise WorkerError("operation_unknown")
    return extract(request)


if __name__ == "__main__":
    try:
        code, payload = 0, json.dumps(main(), ensure_ascii=True, sort_keys=True, allow_nan=False)
    except WorkerError as error:
        code, payload = 3, json.dumps({"status": "error", "code": error.code}, sort_keys=True)
    except Exception as error:  # never echo request content or paths
        code, payload = 4, json.dumps({"status": "error", "code": "worker_failed",
                                       "error_type": type(error).__name__}, sort_keys=True)
    sys.stdout.buffer.write(payload.encode("ascii"))
    sys.stdout.buffer.flush()
    sys.exit(code)
